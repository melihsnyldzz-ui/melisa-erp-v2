from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "erp-v2" / "Melisa_ERP_Personel_Egitimi.docx"

NAVY = "0B2545"
BLUE = "2563EB"
TEAL = "0F766E"
GREEN = "16A34A"
RED = "C9251B"
AMBER = "D97706"
INK = "111827"
MUTED = "64748B"
BORDER = "CBD5E1"
SOFT_BLUE = "EFF6FF"
SOFT_TEAL = "ECFDF5"
SOFT_RED = "FEF2F2"
SOFT_AMBER = "FFFBEB"
SOFT_GRAY = "F8FAFC"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, width in enumerate(widths):
            cell = row.cells[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = tc_pr.first_child_found_in("w:tcBorders")
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                tc_pr.append(borders)
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                tag = f"w:{edge}"
                element = borders.find(qn(tag))
                if element is None:
                    element = OxmlElement(tag)
                    borders.append(element)
                element.set(qn("w:val"), "nil")


def set_run(run, size=10.5, color=INK, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(paragraph, before=0, after=6, line=1.18, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_text(paragraph, text, size=10.5, color=INK, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_run(run, size=size, color=color, bold=bold, italic=italic)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=14 if level == 1 else 10, after=6)
    add_text(
        paragraph,
        text,
        size=16 if level == 1 else 13,
        color=TEAL if level == 1 else NAVY,
        bold=True,
    )
    return paragraph


def add_body(doc, text, after=6):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, after=after)
    add_text(paragraph, text, size=10.5)
    return paragraph


def add_bullet(doc, text, color=INK):
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, after=3, line=1.16)
    if paragraph.runs:
        paragraph.runs[0].text = ""
    add_text(paragraph, text, size=10.2, color=color)


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    style_paragraph(paragraph, after=3, line=1.16)
    if paragraph.runs:
        paragraph.runs[0].text = ""
    add_text(paragraph, text, size=10.2)


def add_callout(doc, title, body, fill=SOFT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=accent, size="8")
    set_cell_margins(cell, top=130, bottom=130, start=180, end=180)
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, after=2)
    add_text(paragraph, title, size=11.2, color=accent, bold=True)
    paragraph = cell.add_paragraph()
    style_paragraph(paragraph, after=0)
    add_text(paragraph, body, size=10.2, color=INK)
    doc.add_paragraph()
    return table


def add_step_box(doc, title, steps, fill=SOFT_GRAY, accent=TEAL):
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, [1680, 7680])
    left, right = table.rows[0].cells
    for cell in (left, right):
        set_cell_border(cell, color=accent, size="8")
        set_cell_margins(cell, top=120, bottom=120, start=140, end=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(left, accent)
    set_cell_shading(right, fill)
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(left.paragraphs[0], after=0)
    add_text(left.paragraphs[0], title, size=11, color=WHITE, bold=True)
    for i, step in enumerate(steps, start=1):
        paragraph = right.paragraphs[0] if i == 1 else right.add_paragraph()
        style_paragraph(paragraph, after=2, line=1.12)
        add_text(paragraph, f"{i}. ", size=10.1, color=accent, bold=True)
        add_text(paragraph, step, size=10.1, color=INK)
    doc.add_paragraph()


def add_flow(doc, items, fills=None):
    fills = fills or [SOFT_BLUE, SOFT_TEAL, SOFT_AMBER, SOFT_RED]
    table = doc.add_table(rows=1, cols=len(items))
    width = int(9360 / len(items))
    set_table_width(table, [width] * len(items))
    for index, item in enumerate(items):
        cell = table.cell(0, index)
        set_cell_shading(cell, fills[index % len(fills)])
        set_cell_border(cell, color=BORDER, size="6")
        set_cell_margins(cell, top=120, bottom=120, start=100, end=100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(paragraph, after=0, line=1.1)
        add_text(paragraph, item, size=9.7, color=INK, bold=True)
    doc.add_paragraph()


def add_matrix(doc, headers, rows, widths, header_fill=TEAL):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, header_fill)
        set_cell_border(cell, color=header_fill, size="8")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        style_paragraph(paragraph, after=0)
        add_text(paragraph, header, size=9.5, color=WHITE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_shading(cell, WHITE if len(table.rows) % 2 else SOFT_GRAY)
            set_cell_border(cell, color=BORDER, size="4")
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            style_paragraph(paragraph, after=0, line=1.12)
            add_text(paragraph, value, size=9.2, color=INK)
    doc.add_paragraph()
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.2)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    style_paragraph(header, after=0)
    add_text(header, "Melisa ERP V2 | Personel Kullanim Egitimi", size=8.8, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(footer, after=0)
    add_text(footer, "Melisa Bebe Tekstil", size=8.5, color=MUTED)

    cover = doc.add_table(rows=1, cols=1)
    set_table_width(cover, [9360])
    cover_cell = cover.cell(0, 0)
    set_cell_shading(cover_cell, SOFT_TEAL)
    set_cell_border(cover_cell, color=TEAL, size="12")
    set_cell_margins(cover_cell, top=420, bottom=420, start=360, end=360)
    p = cover_cell.paragraphs[0]
    style_paragraph(p, after=8)
    add_text(p, "MELISA ERP V2", size=26, color=TEAL, bold=True)
    p = cover_cell.add_paragraph()
    style_paragraph(p, after=10)
    add_text(p, "Personel Kullanim Egitimi", size=18, color=NAVY, bold=True)
    p = cover_cell.add_paragraph()
    style_paragraph(p, after=0)
    add_text(
        p,
        "Satis, stok, cari, tedarikci, tahsilat ve gun sonu kontrollerini basit is akislariyla anlatir.",
        size=11.5,
        color=INK,
    )

    doc.add_paragraph()
    add_flow(doc, ["Urun", "Stok", "Satis / Alis", "Cari hareket", "Ekstre / Rapor"])
    add_callout(
        doc,
        "Egitim hedefi",
        "Personel her isi dogru ekrana yonlendirsin, kaydetmeden once cari/stok/para birimi etkisini kontrol etsin.",
        fill=SOFT_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "1. ERP'nin basit mantigi")
    add_body(
        doc,
        "ERP'de her kayit stok, cari veya kasa tarafinda iz birakir. Bu yuzden ekrana veri girerken sadece formu doldurmus olmayiz; stok miktarini, musteri borcunu veya tedarikci borcunu degistiririz.",
    )
    add_matrix(
        doc,
        ["Islem", "Sistem etkisi", "Kontrol edilecek yer"],
        [
            ("Satis fisi", "Stok azalir, musteri carisine borc yazar.", "Satis Fisleri / Musteri Ekstre"),
            ("Alis fisi", "Stok artar, tedarikci carisine borc yazar.", "Alis Fisleri / Tedarikci Ekstre"),
            ("Tahsilat", "Musteri borcu azalir.", "Tahsilat / Musteri Ekstre"),
            ("Tedarikci odemesi", "Tedarikci borcu azalir.", "Odeme / Tedarikci Ekstre"),
        ],
        [1800, 4300, 3260],
    )

    add_heading(doc, "2. Gunluk acilis rutini")
    add_step_box(
        doc,
        "Acilis",
        [
            "Operasyon Masasi ekranini ac.",
            "Eksik stok uyarilarina bak.",
            "Acik hesap ve tahsilat bekleyenleri kontrol et.",
            "Son satis ve son alis hareketlerinde anormal kayit var mi bak.",
        ],
        fill=SOFT_BLUE,
        accent=BLUE,
    )
    add_callout(
        doc,
        "Kural",
        "Gun basinda once Operasyon Masasi, sonra ilgili detay ekrani acilir. Direkt islem yapmadan once genel durum gorulur.",
        fill=SOFT_AMBER,
        accent=AMBER,
    )

    add_heading(doc, "3. Urun bulma ve stok kontrolu")
    add_step_box(
        doc,
        "Stok",
        [
            "Stok > Stok Listesi ekranini ac.",
            "Stok kodu, barkod, urun adi, marka veya kategori ile ara.",
            "Urunu sec ve sagdaki Urun 360 panelini oku.",
            "Satis, alis, stok, tedarikci ve son hareket bilgilerini kontrol et.",
        ],
        fill=SOFT_TEAL,
        accent=TEAL,
    )
    add_bullet(doc, "Stok eksi ise satis yapmadan once yetkiliye sor.")
    add_bullet(doc, "Fiyat veya tedarikci yanlis gorunuyorsa Urun Karti acilir.")

    add_heading(doc, "4. Yeni satis yapma")
    add_flow(doc, ["Yeni Satis Fisi", "Musteri sec", "Urun ekle", "Adet/fiyat kontrol", "Kaydet"])
    add_step_box(
        doc,
        "Satis",
        [
            "Satis > Yeni Satis Fisi ekranini ac.",
            "Musteriyi sec. Pesin satis ise pesin musteri kullan.",
            "Urunleri barkod, stok kodu veya urun adi ile ekle.",
            "Adet, fiyat, para birimi ve toplam tutari kontrol et.",
            "Kaydet. Kayit sonrasi stok azalir ve musteri carisine borc yazar.",
        ],
        fill=SOFT_BLUE,
        accent=BLUE,
    )
    add_callout(
        doc,
        "Kaydetmeden once",
        "Musteri adi, urun satirlari, adet, fiyat ve para birimi mutlaka kontrol edilir.",
        fill=SOFT_RED,
        accent=RED,
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "5. Musteri cari ve tahsilat")
    add_matrix(
        doc,
        ["Ekran", "Ne icin kullanilir?", "Personelin yapacagi is"],
        [
            ("Musteri Cari Listesi", "Musterileri ve bakiyeleri gormek.", "Bakiyeye gore sirala, cariyi sec."),
            ("Musteri Ekstre", "Satis ve tahsilat hareketlerini birlikte okumak.", "Hangi fis borc yazmis, hangi tahsilat dusmus kontrol et."),
            ("Tahsilat / Odeme", "Musteriden para alindiginda kayit girmek.", "Musteriyi sec, tutar/PB gir, islem etkisini kontrol et."),
        ],
        [2300, 3600, 3460],
    )
    add_step_box(
        doc,
        "Tahsilat",
        [
            "Finans > Tahsilat / Odeme ekranini ac.",
            "Musteri Tahsilati modunu sec.",
            "Musteriyi sec; tarih, tutar, para birimi, odeme sekli ve kasa/banka gir.",
            "Sagdaki Islem Etkisi panelinde kayit sonrasi bakiyeyi kontrol et.",
            "Tahsilat Kaydet ile islemi tamamla.",
        ],
        fill=SOFT_TEAL,
        accent=TEAL,
    )

    add_heading(doc, "6. Tedarikciden mal alis")
    add_flow(doc, ["Alis fisi", "Tedarikci sec", "Urun/adet ekle", "Alis fiyati kontrol", "Kaydet"])
    add_step_box(
        doc,
        "Alis",
        [
            "Alis fisi ekranini ac.",
            "Tedarikciyi sec.",
            "Gelen urunleri barkod veya stok kodu ile ekle.",
            "Adet, alis fiyati, iskonto ve para birimini kontrol et.",
            "Kaydet. Kayit sonrasi stok artar ve tedarikci carisine borc yazar.",
        ],
        fill=SOFT_AMBER,
        accent=AMBER,
    )
    add_bullet(doc, "Yanlis tedarikci secilirse borc yanlis hesaba yazar.", color=RED)
    add_bullet(doc, "Yanlis alis fiyati maliyeti bozar.", color=RED)

    add_heading(doc, "7. Tedarikci cari ve odeme")
    add_step_box(
        doc,
        "Odeme",
        [
            "Cari > Tedarikci Cari Listesi ekraninda tedarikciyi bul.",
            "Bakiyeye gore sirala ve ilgili tedarikciyi sec.",
            "Tedarikci Ekstre ile son alis ve son odeme hareketlerini kontrol et.",
            "Finans > Tahsilat / Odeme ekraninda Tedarikci Odemesi modunu sec.",
            "Tutar, para birimi, kasa/banka ve islem etkisini kontrol edip kaydet.",
        ],
        fill=SOFT_BLUE,
        accent=BLUE,
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "8. Ekstre okuma")
    add_body(
        doc,
        "Ekstre, bir musteri veya tedarikci hesabinin hareket defteridir. Satis, alis, tahsilat, odeme ve iptal hareketleri burada birlikte okunur.",
    )
    add_matrix(
        doc,
        ["Hareket", "Anlami", "Bakiye etkisi"],
        [
            ("Satis fisi", "Musteriye mal kesildi.", "Musteri borcu artar."),
            ("Tahsilat", "Musteriden para alindi.", "Musteri borcu azalir."),
            ("Alis fisi", "Tedarikciden mal geldi.", "Tedarikci borcu artar."),
            ("Odeme", "Tedarikciye para odendi.", "Tedarikci borcu azalir."),
            ("Iptal", "Yanlis/kapatilan kayit.", "Ilgili hareket tersine doner."),
        ],
        [1900, 4000, 3460],
        header_fill=NAVY,
    )

    add_heading(doc, "9. Excel / CSV aktarim kurali")
    add_callout(
        doc,
        "Dikkat",
        "Toplu aktarimdan once yedek alinmali. Yanlis musteri veya tedarikci dosyasi yuklenirse cari listeleri bozulur.",
        fill=SOFT_RED,
        accent=RED,
    )
    add_number(doc, "Aktarim turunu sec: Musteri Cari, Tedarikci Cari veya stok aktarimi.")
    add_number(doc, "Dosyayi yukle ve on izleme ekraninda kolonlari kontrol et.")
    add_number(doc, "Kod, unvan, telefon, para birimi ve bakiye alanlarini kontrol et.")
    add_number(doc, "Yanlis dosya ise kaydetme; dogruysa aktarimi tamamla.")

    add_heading(doc, "10. Gun sonu checklist")
    checks = [
        ("Bugunku satis fisleri dogru mu?", "Satis Fisleri"),
        ("Bugunku alis fisleri dogru mu?", "Alis Fisleri"),
        ("Tahsilatlar dogru cariye islenmis mi?", "Tahsilat / Musteri Ekstre"),
        ("Tedarikci odemeleri dogru mu?", "Odeme / Tedarikci Ekstre"),
        ("Eksi stok var mi?", "Operasyon Masasi / Stok Listesi"),
        ("Kasa/banka hareketi beklenen tutarla uyumlu mu?", "Finans Ozeti"),
    ]
    add_matrix(doc, ["Kontrol", "Bakilacak ekran"], checks, [5200, 4160], header_fill=TEAL)

    add_heading(doc, "11. Personel rolleri")
    add_matrix(
        doc,
        ["Rol", "Ana gorev", "Kullandigi ekranlar"],
        [
            ("Satis", "Urun arama, satis fisi, musteri bilgisini kontrol.", "Stok Listesi, Yeni Satis Fisi, Satis Fisleri"),
            ("Finans", "Tahsilat, odeme, cari ekstre ve kasa kontrolu.", "Tahsilat / Odeme, Musteri Ekstre, Tedarikci Ekstre"),
            ("Depo / Stok", "Stok, alis urun/adet kontrolu, eksi stok takibi.", "Stok Listesi, Alis Fisleri, Stok Hareketleri"),
            ("Yonetici", "Genel durum, raporlar, yetki ve risk kontrolu.", "Operasyon Masasi, Raporlar, Ayarlar"),
        ],
        [1600, 3800, 3960],
        header_fill=NAVY,
    )

    add_heading(doc, "12. En sik yapilan hatalar")
    add_matrix(
        doc,
        ["Hata", "Sonuc", "Dogru davranis"],
        [
            ("Yanlis musteriye satis", "Cari bakiye yanlis olur.", "Kaydetmeden once musteri adini kontrol et."),
            ("Yanlis para birimi", "Bakiye ve rapor bozulur.", "TRY / USD / EUR secimini kontrol et."),
            ("Yanlis adet", "Stok bozulur.", "Satir adetini kontrol et."),
            ("Yanlis tedarikci", "Borc yanlis hesaba yazar.", "Alis fisi basinda tedarikciyi kontrol et."),
            ("Kontrolsuz aktarim", "Cari/stok listeleri bozulur.", "Yedek ve on izleme olmadan aktarim yapma."),
        ],
        [2600, 3100, 3660],
        header_fill=RED,
    )

    add_heading(doc, "13. Kisa egitim uygulamasi")
    add_body(doc, "Yeni personele su 6 uygulama sirayla yaptirilir. Bunlari yapabilen kisi gunluk kullanim mantigini kavramis olur.")
    for item in [
        "Stok Listesi ac, bir urun ara, Urun 360 panelini oku.",
        "Musteri Cari Listesi ac, bakiyesi yuksek musterileri sirala.",
        "Bir musteri ekstresi ac, satis ve tahsilat hareketlerini ayirt et.",
        "Satis Fisleri ekraninda bugunku fisleri filtrele.",
        "Tedarikci Cari Listesi ac, bakiyesi yuksek tedarikcileri sirala.",
        "Tahsilat / Odeme ekraninda kaydetmeden Islem Etkisi panelini oku.",
    ]:
        add_number(doc, item)

    add_callout(
        doc,
        "Son kural",
        "Emin olunmayan kayit kaydedilmez. Once ilgili kart, ekstre veya operasyon masasi kontrol edilir.",
        fill=SOFT_TEAL,
        accent=TEAL,
    )

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
