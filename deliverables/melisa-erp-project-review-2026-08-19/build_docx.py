from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "Melisa_ERP_Proje_Hafizasi_ve_Stratejik_Degerlendirme_2026-08-19.md"
OUT_PATH = ROOT / "Melisa_ERP_Proje_Raporu_2026-08-19.docx"
ASSET_DIR = ROOT / "assets"
ASSET_DIR.mkdir(exist_ok=True)

NAVY = "10233F"
TEAL = "008B8B"
RED = "D64545"
GOLD = "F0A43A"
BLUE = "3D6FE8"
PALE = "EAF0F6"
LIGHT = "F6F8FB"
GREEN = "198754"
GRAY = "5C6B7A"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, field_code: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("MELİSA ERP v2  •  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    add_field(paragraph, "PAGE")


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.extend([color, underline])
    new_run.append(rPr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text: str):
    url_pattern = re.compile(r"https?://\S+")
    pos = 0
    for url_match in url_pattern.finditer(text):
        prefix = text[pos:url_match.start()]
        if prefix:
            add_bold_chunks(paragraph, prefix)
        url = url_match.group(0).rstrip(".,)")
        add_hyperlink(paragraph, url, url)
        pos = url_match.start() + len(url)
    if pos < len(text):
        add_bold_chunks(paragraph, text[pos:])


def add_bold_chunks(paragraph, text: str):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(NAVY)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.2)
            run.font.color.rgb = RGBColor.from_string(TEAL)
        else:
            paragraph.add_run(part)


def _font(size, bold=False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    path = Path("C:/Windows/Fonts") / name
    return ImageFont.truetype(str(path), size)


def _canvas(title, width=1500, height=520):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 28), title, font=_font(30, True), fill="#10233F")
    draw.line((55, 80, width-55, 80), fill="#D8E2EC", width=2)
    return image, draw


def create_charts():
    image, draw = _canvas("Yıllık satış eğilimi – 2026 baz projeksiyondur", height=520)
    datasets = [
        ("USD / milyon $", [3.320, 3.128, 3.020, 2.850], "#3D6FE8"),
        ("EUR / milyon €", [.516, .641, .702, .803], "#008B8B"),
        ("TRY / milyon ₺", [2.967, 2.713, 4.870, 3.870], "#F0A43A"),
    ]
    years = [2023, 2024, 2025, 2026]
    for panel, (label, vals, color) in enumerate(datasets):
        x0 = 55 + panel * 480
        y0, w, h = 120, 410, 295
        draw.text((x0, 92), label, font=_font(19, True), fill="#10233F")
        maxv = max(vals) * 1.18
        for k in range(4):
            gy = y0 + h - k*h/3
            draw.line((x0, gy, x0+w, gy), fill="#E7EDF3", width=1)
        bw = 68
        gap = 30
        for i, (year, val) in enumerate(zip(years, vals)):
            bx = x0 + 28 + i*(bw+gap)
            bh = val/maxv*h
            by = y0+h-bh
            draw.rounded_rectangle((bx, by, bx+bw, y0+h), radius=8, fill=color)
            draw.text((bx+bw/2, by-24), f"{val:.2f}", font=_font(15, True), fill="#27364A", anchor="mm")
            draw.text((bx+bw/2, y0+h+25), str(year), font=_font(14), fill="#5C6B7A", anchor="mm")
    image.save(ASSET_DIR / "sales_trend.png")

    image, draw = _canvas("11.125 aktif stok kartının dağılımı", height=330)
    labels = [("Pozitif", 1843, "#198754"), ("Sıfır", 8930, "#CBD5E1"), ("Negatif", 352, "#D64545")]
    total = sum(v for _, v, _ in labels)
    x, y, width, height = 65, 125, 1370, 72
    cursor = x
    for label, value, color in labels:
        seg = width * value/total
        draw.rectangle((cursor, y, cursor+seg, y+height), fill=color)
        if seg > 120:
            txt_color = "#10233F" if color == "#CBD5E1" else "white"
            draw.text((cursor+seg/2, y+height/2), f"{value/total:.1%}", font=_font(20, True), fill=txt_color, anchor="mm")
        cursor += seg
    lx = 230
    for label, value, color in labels:
        draw.rounded_rectangle((lx, 245, lx+22, 267), radius=4, fill=color)
        draw.text((lx+32, 256), f"{label}: {value:,}".replace(",", "."), font=_font(18), fill="#27364A", anchor="lm")
        lx += 360
    image.save(ASSET_DIR / "stock_health.png")

    image, draw = _canvas("Pazarlama + CRM + operasyon senaryoları", height=580)
    scenarios = [("Temkinli", [100,120,138,155], "#5C6B7A"), ("Merkezi", [100,142,185,222], "#008B8B"), ("Güçlü", [100,170,238,309], "#F0A43A")]
    x0, y0, w, h = 120, 120, 1240, 340
    for k in range(5):
        gy = y0 + k*h/4
        draw.line((x0, gy, x0+w, gy), fill="#E7EDF3", width=1)
    xpts = [x0+i*w/3 for i in range(4)]
    for name, vals, color in scenarios:
        pts = [(x, y0+h-(v/340)*h) for x, v in zip(xpts, vals)]
        draw.line(pts, fill=color, width=7, joint="curve")
        for (x, y), v in zip(pts, vals):
            draw.ellipse((x-8,y-8,x+8,y+8), fill=color)
            draw.text((x, y-22), str(v), font=_font(16, True), fill="#27364A", anchor="mm")
    for x, label in zip(xpts, ["Bugün", "1. yıl", "2. yıl", "3. yıl"]):
        draw.text((x, y0+h+38), label, font=_font(18), fill="#5C6B7A", anchor="mm")
    lx = 430
    for name, vals, color in scenarios:
        draw.line((lx, 525, lx+40, 525), fill=color, width=7)
        draw.text((lx+50, 525), name, font=_font(17), fill="#27364A", anchor="lm")
        lx += 230
    image.save(ASSET_DIR / "growth_scenarios.png")

    image, draw = _canvas("10 milyon USD kontrollü büyüme rotası", height=560)
    years = [2026, 2027, 2028, 2029, 2030]
    values = [3.856, 5.0, 6.5, 8.2, 10.0]
    x0, y0, w, h = 120, 120, 1240, 320
    for k in range(6):
        gy = y0 + k*h/5
        draw.line((x0, gy, x0+w, gy), fill="#E7EDF3", width=1)
    pts = []
    for i, (year, val) in enumerate(zip(years, values)):
        x = x0+i*w/4
        y = y0+h-(val/11)*h
        pts.append((x,y))
        draw.text((x, y0+h+38), str(year), font=_font(18), fill="#5C6B7A", anchor="mm")
    draw.line(pts, fill="#D64545", width=8, joint="curve")
    for (x,y), val in zip(pts, values):
        draw.ellipse((x-10,y-10,x+10,y+10), fill="#D64545")
        draw.text((x, y-28), f"${val:.2f}m", font=_font(18, True), fill="#10233F", anchor="mm")
    draw.text((65, 500), "USD eşdeğeri • Baz ve hedefler yönetim planıdır", font=_font(16), fill="#5C6B7A")
    image.save(ASSET_DIR / "ten_million_roadmap.png")


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string("27364A")
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in [
        ("Title", 34, NAVY, 0, 14),
        ("Heading 1", 20, NAVY, 16, 8),
        ("Heading 2", 14, TEAL, 12, 5),
        ("Heading 3", 11, NAVY, 8, 3),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor.from_string("17324D")
    code.paragraph_format.left_indent = Cm(.35)
    code.paragraph_format.right_indent = Cm(.35)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)

    for style_name in ["List Bullet", "List Number"]:
        styles[style_name].font.name = "Aptos"
        styles[style_name].font.size = Pt(9)
        styles[style_name].paragraph_format.space_after = Pt(2)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("MELİSA BABY  /  YÖNETİM DOSYASI")
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)

    doc.add_paragraph("\n")
    p = doc.add_paragraph()
    r = p.add_run("MELİSA\nERP v2")
    r.font.name = "Aptos Display"
    r.font.size = Pt(42)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    r = p.add_run("Proje Hafızası, Şirket Verisi Özeti\nve Stratejik Değerlendirme")
    r.font.name = "Aptos Display"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    doc.add_paragraph("\n")
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    items = [("244", "commit"), ("43", "ERP Core sprint"), ("8 + 8", "UX ve CRM sprint"), ("5", "canlı servis yüzeyi")]
    for cell, (big, small) in zip(table.rows[0].cells, items):
        set_cell_shading(cell, PALE)
        set_cell_margins(cell, 160, 120, 160, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(big + "\n")
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        r = p.add_run(small)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(GRAY)

    doc.add_paragraph("\n")
    p = doc.add_paragraph()
    r = p.add_run("Rapor tarihi  •  19 Ağustos 2026\nKapsam  •  Mayıs–Ağustos 2026")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("GİZLİLİK NOTU")
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(RED)
    p.add_run("\nŞirket içi kullanım ve kontrollü yapay zekâ değerlendirmesi içindir. Şifre, token ve API anahtarı içermez; ticari veriler içerir.")
    doc.add_page_break()


def add_status_legend(doc: Document):
    doc.add_heading("Doğruluk ve Durum Etiketleri", level=1)
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    items = [
        ("CANLI DOĞRULANDI", "19 Ağustos 2026 tarihinde çalışan servis veya endpoint üzerinden kontrol edildi.", GREEN),
        ("KODDA / BELGEDE VAR", "Kod, test veya dokümanda mevcut; tüm saha senaryoları tamamlanmış sayılmaz.", BLUE),
        ("TARİHLİ VERİ", "Salt okunur şirket verisi anlık görüntüsü; tarih ilerledikçe değişir.", GOLD),
        ("PLAN / HEDEF", "Konuşuldu veya tasarlandı; tamamlanmış özellik değildir.", GRAY),
    ]
    for row, (name, desc, color) in zip(table.rows, items):
        set_cell_shading(row.cells[0], color)
        set_cell_shading(row.cells[1], LIGHT)
        for cell in row.cells:
            set_cell_margins(cell)
        p = row.cells[0].paragraphs[0]
        r = p.add_run(name)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        row.cells[1].text = desc
    doc.add_paragraph("Bu ayrım, projenin başarısını küçültmeden mevcut üretim risklerini dürüst biçimde korur.")
    doc.add_page_break()


def add_toc(doc: Document):
    doc.add_heading("İçindekiler", level=1)
    sections = [
        "1. Yönetici Özeti", "2. Şirket Vizyonu ve İş Kuralları",
        "3. Projenin Kronolojik Hikâyesi", "4. Sistem Mimarisi",
        "5. Başardıklarımız: Modül Modül", "6. Test, Kabul ve Kalite Kanıtları",
        "7. Vega 2026 Canlı ve Salt Okunur Entegrasyon",
        "8. Şirket Verilerinden Elde Edilen Bulgular", "9. CRM ve Ticari Dönüşüm Hedefi",
        "10. Büyüme Senaryoları", "11. Pazarlama ve Pazar Yaklaşımı",
        "12. Açık Riskler ve Tamamlanması Gerekenler", "13. Önerilen Önceliklendirme",
        "14. Yönetim KPI Seti", "15. Başka Bir Yapay Zekâ İçin Değerlendirme Talimatı",
        "16. Kaynak ve Dayanak Özeti", "17. Sonuç",
    ]
    table = doc.add_table(rows=9, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, title in enumerate(sections):
        cell = table.cell(idx % 9, idx // 9)
        set_cell_margins(cell, 65, 80, 65, 80)
        if idx % 2 == 0:
            set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(NAVY)
        if idx in (0, 7, 11):
            r.font.bold = True
    p = doc.add_paragraph("Word sürümünde başlıklar gezinme panelinden de izlenebilir.")
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_page_break()


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_md_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                set_cell_shading(cell, NAVY)
            elif ri % 2 == 0:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if ci > 0 and re.search(r"\d", row[ci] if ci < len(row) else "") else WD_ALIGN_PARAGRAPH.LEFT
            add_rich_text(p, row[ci] if ci < len(row) else "")
            for run in p.runs:
                run.font.size = Pt(7.8)
                if ri == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(tbl_header)
    doc.add_paragraph()


def insert_chart(doc, filename, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSET_DIR / filename), width=Inches(6.65))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_markdown_body(doc: Document):
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    # Skip source cover and front-matter before section 1; custom front pages already cover them.
    start = next(i for i, line in enumerate(lines) if line.startswith("# 1. "))
    lines = lines[start:]
    i = 0
    code_mode = False
    code_buffer = []
    quote_buffer = []
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if code_mode:
                p = doc.add_paragraph("\n".join(code_buffer), style="Code Block")
                set_cell = OxmlElement("w:shd")
                set_cell.set(qn("w:fill"), PALE)
                p._p.get_or_add_pPr().append(set_cell)
                code_buffer = []
                code_mode = False
            else:
                code_mode = True
            i += 1
            continue
        if code_mode:
            code_buffer.append(line)
            i += 1
            continue
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i+1].strip().startswith("|"):
            rows, i = parse_table(lines, i)
            add_md_table(doc, rows)
            continue
        if line.startswith("# "):
            if not first_h1:
                doc.add_page_break()
            title = line[2:].strip()
            doc.add_heading(title, level=1)
            first_h1 = False
            if title.startswith("8. Şirket Verilerinden"):
                pass
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            doc.add_heading(title, level=2)
            if title.startswith("8.1 Satış"):
                insert_chart(doc, "sales_trend.png", "Kaynak: Vega salt okunur satış görüntüsü; 2026 değerleri baz projeksiyondur.")
            elif title.startswith("8.2 Stok"):
                insert_chart(doc, "stock_health.png", "Aktif kartların büyük çoğunluğu sıfır stoktadır; negatif stok ayrı kontrol gerektirir.")
            elif title.startswith("10.1 Pazarlama"):
                insert_chart(doc, "growth_scenarios.png", "Senaryolar tahmindir; stok, insan ve çalışma sermayesi kapasitesine bağlıdır.")
            elif title.startswith("10.2 10 milyon"):
                insert_chart(doc, "ten_million_roadmap.png", "Baz yaklaşık 3,856 milyon USD eşdeğeridir; hedef kontrollü yıllık artışla 2030’dur.")
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(.45)
            p.paragraph_format.first_line_indent = Cm(-.45)
            add_rich_text(p, stripped)
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, stripped[2:])
            i += 1
            continue
        if stripped.startswith(">"):
            q = stripped.lstrip("> ")
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(.55)
            p.paragraph_format.right_indent = Cm(.35)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_rich_text(p, q)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(TEAL)
            i += 1
            continue

        p = doc.add_paragraph()
        add_rich_text(p, stripped)
        i += 1


def main():
    create_charts()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(.6)
    section.footer_distance = Cm(.6)
    configure_styles(doc)
    add_cover(doc)
    add_status_legend(doc)
    add_toc(doc)
    add_markdown_body(doc)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "MELİSA ERP v2  •  PROJE HAFIZASI VE STRATEJİK DEĞERLENDİRME"
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in header.runs:
            run.font.size = Pt(7.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(TEAL)
        add_page_number(section.footer.paragraphs[0])

    props = doc.core_properties
    props.title = "Melisa ERP v2 Proje Hafızası ve Stratejik Değerlendirme"
    props.subject = "Mayıs–Ağustos 2026 proje özeti, şirket verisi ve büyüme planı"
    props.author = "Melisa Baby / Codex"
    props.keywords = "ERP, CRM, Vega, WhatsApp, stok, satış, strateji"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
